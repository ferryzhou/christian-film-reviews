#!/usr/bin/env python3
"""把 original-reviews/*.md 编成一本可上架 Amazon KDP 的电子书。

用法（在仓库任意位置运行）：
  python3 book/build_book.py                 # 生成 book/dist/ 下全部纯文字版本 + 封面
  python3 book/build_book.py --with-images   # 另外生成内嵌剧照的版本到 book/dist/illustrated/（版权风险见 book/README.md）
  python3 book/build_book.py --only docx-tw  # 只生成某一种：docx-tw | docx-sc | epub-tw | epub-sc | cover

产出：
  dist/<书名>-繁體.docx   Amazon KDP 上传用（KDP 的中文电子书只接受繁体中文 + DOCX，见 README）
  dist/<书名>-简体.docx   备用（Word 校对 / 其他平台）
  dist/<书名>-繁體.epub   Kindle Send-to-Kindle、Apple Books、Google Play、Kobo 等
  dist/<书名>-简体.epub   同上，简体
  dist/cover-tw.jpg / cover-sc.jpg   1600×2560 电子书封面（原创排版，不含任何影片素材）
  dist/kdp-listing-tw.txt / -sc.txt  上架页面用的书名/简介/关键词，直接复制粘贴

依赖：pip install python-docx ebooklib opencc-python-reimplemented
封面另需 node + playwright（本仓库配套 render_cover.js）；缺失时跳过封面并提示。
"""

import argparse
import json
import os
import re
import subprocess
import sys
import uuid

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT)

from publish_wordpress import load_registry, parse_front_matter  # noqa: E402
from build_review_pages import load_films  # noqa: E402

CONFIG = json.load(open(os.path.join(HERE, "book.json"), encoding="utf-8"))
DIST = os.path.join(HERE, "dist")

# 圣经 66 卷正典次序（和合本卷名），用于经文索引排序
BIBLE_BOOKS = [
    "创世记", "出埃及记", "利未记", "民数记", "申命记", "约书亚记", "士师记", "路得记",
    "撒母耳记上", "撒母耳记下", "列王纪上", "列王纪下", "历代志上", "历代志下", "以斯拉记",
    "尼希米记", "以斯帖记", "约伯记", "诗篇", "箴言", "传道书", "雅歌", "以赛亚书",
    "耶利米书", "耶利米哀歌", "以西结书", "但以理书", "何西阿书", "约珥书", "阿摩司书",
    "俄巴底亚书", "约拿书", "弥迦书", "那鸿书", "哈巴谷书", "西番雅书", "哈该书",
    "撒迦利亚书", "玛拉基书",
    "马太福音", "马可福音", "路加福音", "约翰福音", "使徒行传", "罗马书", "哥林多前书",
    "哥林多后书", "加拉太书", "以弗所书", "腓立比书", "歌罗西书", "帖撒罗尼迦前书",
    "帖撒罗尼迦后书", "提摩太前书", "提摩太后书", "提多书", "腓利门书", "希伯来书",
    "雅各书", "彼得前书", "彼得后书", "约翰一书", "约翰二书", "约翰三书", "犹大书", "启示录",
]
SCRIPTURE_RE = re.compile(
    r"[（(]([一-鿿]+)\s*(\d+:\d+(?:[–\-—]\d+(?::\d+)?)?(?:[,，、]\s*\d+(?::\d+)?)*)[)）]"
)
CN_NUM = "一二三四五六七八九十"


# ---------------------------------------------------------------- 文本与语言

class Lang:
    """简体 / 繁体两个版本共用一套内容，差别只在 T() 的转换与少数标签。"""

    def __init__(self, code):
        self.code = code  # "sc" | "tw"
        if code == "tw":
            from opencc import OpenCC
            self._cc = OpenCC("s2tw")
            self.bcp47, self.label, self.file_tag = "zh-TW", "繁體中文", "繁體"
        else:
            self._cc = None
            self.bcp47, self.label, self.file_tag = "zh-CN", "简体中文", "简体"

    def T(self, s):
        return self._cc.convert(s) if self._cc else s


# ---------------------------------------------------------------- 内容装配

def parse_blocks(md):
    """本站 Markdown 子集 -> [("p", text) | ("img", caption, relpath)]（忽略一级标题）。"""
    blocks = []
    for b in (x.strip() for x in re.split(r"\n{2,}", md)):
        if not b or b.startswith("# "):
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)\s]+)\)$", b)
        if m:
            blocks.append(("img", m.group(1), m.group(2)))
        elif b.startswith("## "):
            blocks.append(("h", b[3:]))
        else:
            blocks.append(("p", re.sub(r"\s*\n\s*", "", b)))
    return blocks


def inline_runs(text):
    """**粗体** / *斜体* -> [(text, bold, italic)]"""
    runs = []
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            runs.append((tok[2:-2], True, False))
        elif tok.startswith("*"):
            runs.append((tok[1:-1], False, True))
        else:
            runs.append((tok, False, False))
    return runs


def load_book():
    """按 book.json 的分辑次序装配全部章节。"""
    registry = load_registry()
    films = load_films()
    parts, seen = [], set()
    for i, part in enumerate(CONFIG["parts"]):
        chapters = []
        for fid in part["films"]:
            if fid not in registry:
                sys.exit(f"book.json 引用了 originals.js 中不存在的影评：{fid}")
            seen.add(fid)
            chapters.append(load_chapter(fid, registry[fid], films.get(fid, {})))
        parts.append({"no": i + 1, "title": part["title"], "theme": part["theme"], "chapters": chapters})
    missing = [f for f in registry if f not in seen]
    if missing:  # 新写的影评尚未分辑：归入末尾"未分辑"，并提醒维护 book.json
        print(f"!! 以下影评未在 book.json 分辑中，暂归入第 {len(parts) + 1} 辑：{', '.join(missing)}")
        parts.append({"no": len(parts) + 1, "title": "余音", "theme": "未分辑",
                      "chapters": [load_chapter(f, registry[f], films.get(f, {})) for f in missing]})
    n = 0
    for part in parts:
        for ch in part["chapters"]:
            n += 1
            ch["no"] = n
            ch["anchor"] = f"ch{n:02d}"
    return parts


def load_chapter(fid, meta, film):
    path = os.path.join(ROOT, "original-reviews", f"{fid}.md")
    fm, body = parse_front_matter(open(path, encoding="utf-8").read())
    blocks = parse_blocks(body)
    text = " ".join(b[1] for b in blocks if b[0] == "p")
    refs = [(m.group(1), m.group(2)) for m in SCRIPTURE_RE.finditer(text)]
    return {
        "id": fid, "title": meta["title"],
        "film": fm.get("film") or film.get("title", ""), "filmEn": fm.get("filmEn") or film.get("titleEn", ""),
        "year": fm.get("year") or film.get("year", ""), "director": fm.get("director") or film.get("director", ""),
        "country": film.get("country", ""), "date": meta.get("date", ""),
        "blocks": blocks, "refs": refs,
    }


def chapter_subtitle(ch, lang):
    """《片名》（English · 年份 · 导演），繁体版附台湾译名。"""
    s = f"《{lang.T(ch['film'])}》"
    tw = CONFIG.get("titleTW", {}).get(ch["id"])
    if lang.code == "tw" and tw and tw != lang.T(ch["film"]):
        s += f"（台譯《{tw}》）"
    bits = [ch["filmEn"], str(ch["year"]), lang.T("导演：") + lang.T(ch["director"])]
    return s + "　" + " · ".join(b for b in bits if b)


def scripture_index(parts):
    """卷名 -> [(章节, 章节号, 经节)]，按正典次序。"""
    idx = {}
    for part in parts:
        for ch in part["chapters"]:
            for book, ref in ch["refs"]:
                idx.setdefault(book, []).append((ch, ref))
    order = {b: i for i, b in enumerate(BIBLE_BOOKS)}

    def ref_key(item):  # 同一卷内按 章:节 数字排序
        nums = re.findall(r"\d+", item[1])
        return (int(nums[0]), int(nums[1]) if len(nums) > 1 else 0)

    return sorted(((bk, sorted(items, key=ref_key)) for bk, items in idx.items()),
                  key=lambda kv: order.get(kv[0], 999))


def part_heading(part, lang):
    no = CN_NUM[part["no"] - 1] if part["no"] <= 10 else str(part["no"])
    return lang.T(f"第{no}辑　{part['title']}"), lang.T(part["theme"])


def copyright_lines(lang):
    c = CONFIG
    return [lang.T(x) for x in [
        f"{c['title']}：{c['subtitle']}",
        f"作者：{c['author']}",
        f"© {c['year']} {c['author']}　保留所有权利。",
        f"{c['edition']}　{c['year']} 年",
        "本书文字均为原创评论。所评影片及其片名、台词、剧照之版权归各出品方与发行方所有，本书引用仅为评论与研究目的。",
        "圣经引文出自《圣经》和合本。",
        f"文章网络版：{c['site']}",
        "封面设计：" + c["publisher"],
    ]]


def preface_blocks():
    md = open(os.path.join(HERE, "preface.md"), encoding="utf-8").read()
    title = re.search(r"^# (.+)$", md, re.M).group(1)
    return title, parse_blocks(md)


def out_name(lang, ext):
    return os.path.join(DIST, f"{CONFIG['title']}-{lang.file_tag}.{ext}")


# ---------------------------------------------------------------- DOCX（KDP 上传格式）

def build_docx(parts, lang, with_images):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches, RGBColor

    T = lang.T
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(6), Inches(9)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.8))

    east_font = "Noto Serif CJK TC" if lang.code == "tw" else "Noto Serif CJK SC"

    def set_font(style, size=None, bold=None, color=None):
        style.font.name = "Noto Serif CJK SC"
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), east_font)
        if size:
            style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if color:
            style.font.color.rgb = RGBColor.from_string(color)

    normal = doc.styles["Normal"]
    set_font(normal, 11)
    normal.paragraph_format.first_line_indent = Pt(22)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5
    set_font(doc.styles["Heading 1"], 20, True, "2A2418")
    set_font(doc.styles["Heading 2"], 16, True, "2A2418")
    set_font(doc.styles["Heading 3"], 12, False, "5A5040")
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].paragraph_format.first_line_indent = Pt(0)
        doc.styles[name].paragraph_format.space_before = Pt(18 if name != "Heading 3" else 0)
        doc.styles[name].paragraph_format.space_after = Pt(12)

    def para(text="", style=None, align=None, indent=True, size=None, italic=False, color=None):
        p = doc.add_paragraph(style=style)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not indent:
            p.paragraph_format.first_line_indent = Pt(0)
        if text:
            for t, b, i in inline_runs(text):
                r = p.add_run(t)
                r.bold, r.italic = b, (i or italic)
                if size:
                    r.font.size = Pt(size)
                if color:
                    r.font.color.rgb = RGBColor.from_string(color)
        return p

    def page_break():
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    _bm_id = [0]

    def bookmark(paragraph, name):
        _bm_id[0] += 1
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), str(_bm_id[0])); start.set(qn("w:name"), name)
        end.set(qn("w:id"), str(_bm_id[0]))
        ppr = paragraph._p.pPr  # bookmarkStart 必须排在 pPr 之后，否则不合 OOXML 规范（Word/LibreOffice 拒开）
        paragraph._p.insert(list(paragraph._p).index(ppr) + 1 if ppr is not None else 0, start)
        paragraph._p.append(end)

    def toc_link(paragraph, text, anchor, indent_pt=0, size=None):
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(indent_pt)
        paragraph.paragraph_format.space_after = Pt(2)
        h = OxmlElement("w:hyperlink")
        h.set(qn("w:anchor"), anchor)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        if size:
            sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(size * 2)); rpr.append(sz)
        r.append(rpr)
        t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
        r.append(t)
        h.append(r)
        paragraph._p.append(h)

    # 书名页
    for _ in range(6):
        para()
    para(T(CONFIG["title"]), align="center", indent=False, size=30)
    para(T(CONFIG["subtitle"]), align="center", indent=False, size=15, color="5A5040")
    for _ in range(3):
        para()
    para(T(CONFIG["author"]), align="center", indent=False, size=13)
    page_break()

    # 版权页
    for line in copyright_lines(lang):
        para(line, indent=False, size=9, color="5A5040")
    page_break()

    # 目录（Kindle 用书签 "toc" 识别逻辑目录）
    h = doc.add_heading(T("目录"), level=1)
    bookmark(h, "toc")
    pre_title, pre_blocks = preface_blocks()
    toc_link(doc.add_paragraph(), T(pre_title), "preface")
    for part in parts:
        ptitle, theme = part_heading(part, lang)
        toc_link(doc.add_paragraph(), f"{ptitle}　{theme}", f"part{part['no']}")
        for ch in part["chapters"]:
            toc_link(doc.add_paragraph(), f"{T(ch['title'])}　《{T(ch['film'])}》", ch["anchor"], indent_pt=18, size=10)
    toc_link(doc.add_paragraph(), T("附录一　影片索引"), "app-films")
    toc_link(doc.add_paragraph(), T("附录二　经文索引"), "app-scripture")
    toc_link(doc.add_paragraph(), T("关于作者"), "about")
    page_break()

    # 序
    h = doc.add_heading(T(pre_title), level=1)
    bookmark(h, "preface")
    for b in pre_blocks:
        if b[0] == "p":
            para(T(b[1]))
    page_break()

    # 正文
    for part in parts:
        ptitle, theme = part_heading(part, lang)
        for _ in range(8):
            para()
        h = doc.add_heading(ptitle, level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        bookmark(h, f"part{part['no']}")
        para(theme, align="center", indent=False, size=13, color="5A5040")
        page_break()
        for ch in part["chapters"]:
            h = doc.add_heading(T(ch["title"]), level=2)
            bookmark(h, ch["anchor"])
            para(chapter_subtitle(ch, lang), style="Heading 3")
            for b in ch["blocks"]:
                if b[0] == "p":
                    para(T(b[1]))
                elif b[0] == "h":
                    para(T(b[1]), style="Heading 3")
                elif b[0] == "img" and with_images:
                    fp = os.path.join(ROOT, b[2])
                    if os.path.exists(fp):
                        pic = doc.add_paragraph()
                        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        pic.paragraph_format.first_line_indent = Pt(0)
                        pic.add_run().add_picture(fp, width=Inches(4.2))
                        para(T(b[1]), align="center", indent=False, size=9, italic=True, color="5A5040")
            page_break()

    # 附录一：影片索引
    h = doc.add_heading(T("附录一　影片索引"), level=1)
    bookmark(h, "app-films")
    rows = [ch for part in parts for ch in part["chapters"]]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light List"
    for i, head in enumerate([T("片名"), T("原名"), T("年份 · 导演"), T("本书篇目")]):
        table.rows[0].cells[i].text = head
    for ch in sorted(rows, key=lambda c: int(c["year"] or 0)):
        cells = table.add_row().cells
        cells[0].text = T(ch["film"]); cells[1].text = ch["filmEn"]
        cells[2].text = f"{ch['year']} · {T(ch['director'])}"; cells[3].text = T(ch["title"])
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    r.font.size = Pt(9)
    page_break()

    # 附录二：经文索引
    h = doc.add_heading(T("附录二　经文索引"), level=1)
    bookmark(h, "app-scripture")
    para(T("以下为各篇正文中引用的和合本经文，按圣经卷序排列；括号内为所在篇目。"), indent=False, size=10, color="5A5040")
    for book, items in scripture_index(parts):
        p = para(indent=False)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(T(book)).bold = True
        p.add_run("　" + "；".join(f"{ref}（{T(ch['title'])}）" for ch, ref in items)).font.size = Pt(10)
    page_break()

    # 关于作者
    h = doc.add_heading(T("关于作者"), level=1)
    bookmark(h, "about")
    para(T(CONFIG["authorBio"]))
    para(CONFIG["site"], indent=False, size=10, color="9A7320")

    doc.core_properties.title = T(f"{CONFIG['title']}：{CONFIG['subtitle']}")
    doc.core_properties.author = T(CONFIG["author"])
    doc.core_properties.language = lang.bcp47
    doc.core_properties.subject = T("电影评论 · 基督信仰")

    out = out_name(lang, "docx")
    doc.save(out)
    return out


# ---------------------------------------------------------------- EPUB（其他平台 / Send to Kindle）

EPUB_CSS = """
@charset "utf-8";
body { font-family: "Noto Serif CJK SC", "Noto Serif SC", "Songti SC", "SimSun", serif; line-height: 1.75; margin: 0 4%; color: #2a2418; }
h1, h2, h3 { font-weight: 600; line-height: 1.3; text-align: left; }
h1 { font-size: 1.7em; margin: 1.6em 0 0.8em; }
h2 { font-size: 1.4em; margin: 1.4em 0 0.3em; }
h3.film { font-size: 0.95em; font-weight: 400; color: #5a5040; margin: 0 0 1.6em; }
p { text-indent: 2em; margin: 0 0 0.7em; text-align: justify; }
p.noindent, p.center, figcaption, .copyright p, .toc p { text-indent: 0; }
.center { text-align: center; }
.titlepage { text-align: center; margin-top: 30%; }
.titlepage .title { font-size: 2.4em; margin: 0; }
.titlepage .subtitle { font-size: 1.2em; color: #5a5040; margin: 0.6em 0 3em; }
.titlepage .author { font-size: 1.1em; }
.copyright p { font-size: 0.85em; color: #5a5040; margin-bottom: 0.5em; }
.part { text-align: center; margin-top: 35%; }
.part h1 { text-align: center; }
.part .theme { color: #5a5040; font-size: 1.1em; text-indent: 0; }
figure { margin: 1.4em 0; text-align: center; }
figure img { max-width: 100%; height: auto; }
figcaption { font-size: 0.82em; color: #5a5040; margin-top: 0.4em; text-align: center; }
table { border-collapse: collapse; width: 100%; font-size: 0.85em; }
th, td { border-bottom: 1px solid #ddd; padding: 0.4em 0.3em; text-align: left; vertical-align: top; }
.scripture p { text-indent: 0; margin-bottom: 0.5em; }
.toc p { margin: 0.25em 0; }
.toc p.sub { margin-left: 1.5em; font-size: 0.95em; }
a { color: #9a7320; text-decoration: none; }
"""


def build_epub(parts, lang, with_images, cover_path):
    import html as H
    from ebooklib import epub

    T = lang.T

    def esc(s):
        return H.escape(s, quote=False)

    def inline(text):
        out = ""
        for t, b, i in inline_runs(text):
            t = esc(t)
            if b:
                t = f"<strong>{t}</strong>"
            if i:
                t = f"<em>{t}</em>"
            out += t
        return out

    book = epub.EpubBook()
    book.set_identifier("urn:uuid:" + str(uuid.uuid5(uuid.NAMESPACE_URL, f"{CONFIG['site']}#book-{lang.code}-{CONFIG['year']}")))
    book.set_title(T(f"{CONFIG['title']}：{CONFIG['subtitle']}"))
    book.set_language(lang.bcp47)
    book.add_author(T(CONFIG["author"]))
    book.add_metadata("DC", "publisher", T(CONFIG["publisher"]))
    book.add_metadata("DC", "description", T(CONFIG["description"]).replace("\n", " "))
    book.add_metadata("DC", "date", str(CONFIG["year"]))
    if cover_path and os.path.exists(cover_path):
        book.set_cover("cover.jpg", open(cover_path, "rb").read())

    css = epub.EpubItem(uid="style", file_name="style/book.css", media_type="text/css", content=EPUB_CSS.encode())
    book.add_item(css)

    def page(uid, fname, title, body_html):
        c = epub.EpubHtml(uid=uid, file_name=fname, title=title, lang=lang.bcp47)
        c.content = f'<html xmlns="http://www.w3.org/1999/xhtml"><head><title>{esc(title)}</title></head><body>{body_html}</body></html>'
        c.add_item(css)
        book.add_item(c)
        return c

    spine, toc = ["nav"], []

    title_html = (f'<div class="titlepage"><p class="title noindent">{esc(T(CONFIG["title"]))}</p>'
                  f'<p class="subtitle noindent">{esc(T(CONFIG["subtitle"]))}</p>'
                  f'<p class="author noindent">{esc(T(CONFIG["author"]))}</p></div>')
    spine.append(page("titlepage", "titlepage.xhtml", T("书名页"), title_html))

    cr_html = '<div class="copyright">' + "".join(f"<p>{esc(x)}</p>" for x in copyright_lines(lang)) + "</div>"
    spine.append(page("copyright", "copyright.xhtml", T("版权页"), cr_html))

    pre_title, pre_blocks = preface_blocks()
    pre_html = f"<h1>{esc(T(pre_title))}</h1>" + "".join(f"<p>{inline(T(b[1]))}</p>" for b in pre_blocks if b[0] == "p")
    c = page("preface", "preface.xhtml", T(pre_title), pre_html)
    spine.append(c); toc.append(c)

    img_items = {}

    def img_item(rel):
        if rel in img_items:
            return img_items[rel]
        fp = os.path.join(ROOT, rel)
        if not os.path.exists(fp):
            return None
        ext = os.path.splitext(rel)[1].lower().lstrip(".")
        fname = "images/" + rel.replace("/", "_")
        it = epub.EpubItem(uid=f"img{len(img_items)}", file_name=fname,
                           media_type="image/png" if ext == "png" else "image/jpeg", content=open(fp, "rb").read())
        book.add_item(it)
        img_items[rel] = fname
        return fname

    for part in parts:
        ptitle, theme = part_heading(part, lang)
        phtml = f'<div class="part"><h1>{esc(ptitle)}</h1><p class="theme">{esc(theme)}</p></div>'
        pc = page(f"part{part['no']}", f"part{part['no']}.xhtml", ptitle, phtml)
        spine.append(pc)
        sub = []
        for ch in part["chapters"]:
            body = f"<h2>{esc(T(ch['title']))}</h2><h3 class=\"film\">{esc(chapter_subtitle(ch, lang))}</h3>"
            for b in ch["blocks"]:
                if b[0] == "p":
                    body += f"<p>{inline(T(b[1]))}</p>"
                elif b[0] == "h":
                    body += f"<h3>{inline(T(b[1]))}</h3>"
                elif b[0] == "img" and with_images:
                    fname = img_item(b[2])
                    if fname:
                        body += f'<figure><img src="{fname}" alt="{H.escape(T(b[1]), quote=True)}"/><figcaption>{inline(T(b[1]))}</figcaption></figure>'
            cc = page(ch["anchor"], f"{ch['anchor']}.xhtml", T(ch["title"]), body)
            spine.append(cc); sub.append(cc)
        toc.append((epub.Section(ptitle, href=pc.file_name), sub))

    rows = sorted((ch for part in parts for ch in part["chapters"]), key=lambda c: int(c["year"] or 0))
    films_html = f"<h1>{esc(T('附录一　影片索引'))}</h1><table><tr><th>{esc(T('片名'))}</th><th>{esc(T('原名'))}</th><th>{esc(T('年份 · 导演'))}</th><th>{esc(T('本书篇目'))}</th></tr>"
    for ch in rows:
        films_html += (f"<tr><td>{esc(T(ch['film']))}</td><td>{esc(ch['filmEn'])}</td>"
                       f"<td>{esc(str(ch['year']))} · {esc(T(ch['director']))}</td>"
                       f"<td><a href=\"{ch['anchor']}.xhtml\">{esc(T(ch['title']))}</a></td></tr>")
    films_html += "</table>"
    c = page("app-films", "app-films.xhtml", T("附录一　影片索引"), films_html)
    spine.append(c); toc.append(c)

    sc_html = f"<h1>{esc(T('附录二　经文索引'))}</h1><p class=\"noindent\">{esc(T('以下为各篇正文中引用的和合本经文，按圣经卷序排列；括号内为所在篇目。'))}</p><div class=\"scripture\">"
    for bk, items in scripture_index(parts):
        refs = "；".join(f"{esc(ref)}（<a href=\"{ch['anchor']}.xhtml\">{esc(T(ch['title']))}</a>）" for ch, ref in items)
        sc_html += f"<p><strong>{esc(T(bk))}</strong>　{refs}</p>"
    sc_html += "</div>"
    c = page("app-scripture", "app-scripture.xhtml", T("附录二　经文索引"), sc_html)
    spine.append(c); toc.append(c)

    about_html = f"<h1>{esc(T('关于作者'))}</h1><p>{esc(T(CONFIG['authorBio']))}</p><p class=\"noindent\"><a href=\"{CONFIG['site']}\">{CONFIG['site']}</a></p>"
    c = page("about", "about.xhtml", T("关于作者"), about_html)
    spine.append(c); toc.append(c)

    book.toc = toc
    book.spine = spine
    book.add_item(epub.EpubNcx())
    nav = epub.EpubNav()
    nav.add_item(css)
    book.add_item(nav)

    out = out_name(lang, "epub")
    epub.write_epub(out, book, {})
    return out


# ---------------------------------------------------------------- 封面与上架文案

def build_cover(lang):
    out = os.path.join(DIST, f"cover-{lang.code}.jpg")
    env = dict(os.environ)
    try:
        node_root = subprocess.check_output(["npm", "root", "-g"], text=True).strip()
        env["NODE_PATH"] = node_root + (os.pathsep + env["NODE_PATH"] if env.get("NODE_PATH") else "")
    except Exception:
        pass
    payload = json.dumps({
        "title": lang.T(CONFIG["title"]), "subtitle": lang.T(CONFIG["subtitle"]),
        "author": lang.T(CONFIG["author"]), "tagline": lang.T("二十一篇以基督信仰为眼光的电影随笔"),
        "out": out,
    }, ensure_ascii=False)
    try:
        subprocess.run(["node", os.path.join(HERE, "render_cover.js"), payload], check=True, env=env)
    except (subprocess.CalledProcessError, FileNotFoundError) as e:
        print(f"!! 封面未生成（需要 node + playwright + Chromium）：{e}")
        return None
    return out


def write_listing(lang):
    T = lang.T
    out = os.path.join(DIST, f"kdp-listing-{lang.code}.txt")
    lines = [
        "=== Kindle eBook Details（复制到 KDP 表单）===", "",
        f"Language: {'Chinese (Traditional)' if lang.code == 'tw' else 'Chinese (Simplified) —— KDP 不支持，此版本不能上传 KDP'}",
        f"Book Title: {T(CONFIG['title'])}",
        f"Subtitle: {T(CONFIG['subtitle'])}",
        f"Author: {T(CONFIG['author'])}", "",
        "Description:", T(CONFIG["description"]), "",
        "Keywords（7 组，每组一格）:",
        *[f"  {i + 1}. {T(k)}" for i, k in enumerate(CONFIG["keywords"])], "",
        "Categories（选 3 个）:", *[f"  - {c}" for c in CONFIG["categories"]], "",
        "Publishing Rights: I own the copyright and I hold the necessary publishing rights",
        "Primary Audience: Not sexually explicit; Reading age: 不填", "",
        "=== 章节一览 ===",
    ]
    for part in load_book():
        ptitle, theme = part_heading(part, lang)
        lines.append(f"{ptitle}　{theme}")
        for ch in part["chapters"]:
            lines.append(f"  {ch['no']:>2}. {T(ch['title'])}　《{T(ch['film'])}》 {ch['filmEn']} ({ch['year']})")
    open(out, "w", encoding="utf-8").write("\n".join(lines) + "\n")
    return out


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--with-images", action="store_true", help="另生成内嵌剧照版到 dist/illustrated/")
    ap.add_argument("--only", choices=["docx-tw", "docx-sc", "epub-tw", "epub-sc", "cover"], help="只生成某一项")
    args = ap.parse_args()
    global DIST

    base = DIST
    os.makedirs(base, exist_ok=True)
    parts = load_book()
    n_ch = sum(len(p["chapters"]) for p in parts)
    print(f"共 {len(parts)} 辑 {n_ch} 篇，{sum(len(c['refs']) for p in parts for c in p['chapters'])} 处经文引用")

    langs = [Lang("tw"), Lang("sc")]
    covers = {}
    if args.only in (None, "cover"):
        for lang in langs:
            covers[lang.code] = build_cover(lang)
            if covers[lang.code]:
                print("封面  ->", covers[lang.code])
    else:
        covers = {l.code: os.path.join(base, f"cover-{l.code}.jpg") for l in langs}

    variants = [(False, base)] + ([(True, os.path.join(base, "illustrated"))] if args.with_images else [])
    for with_images, outdir in variants:
        DIST = outdir
        os.makedirs(DIST, exist_ok=True)
        for lang in langs:
            if args.only in (None, f"docx-{lang.code}"):
                print("DOCX  ->", build_docx(parts, lang, with_images))
            if args.only in (None, f"epub-{lang.code}"):
                print("EPUB  ->", build_epub(parts, lang, with_images, covers.get(lang.code)))
    DIST = base
    if args.only is None:
        for lang in langs:
            print("上架文案 ->", write_listing(lang))


if __name__ == "__main__":
    main()
